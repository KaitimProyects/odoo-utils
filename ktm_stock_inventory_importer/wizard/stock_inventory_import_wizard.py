import base64
import binascii
from collections import defaultdict
from io import BytesIO

from openpyxl import load_workbook

from odoo import _, api, fields, models
from odoo.exceptions import UserError

# Canonical column key -> accepted headers (compared trimmed and casefolded).
HEADER_ALIASES = {
    'name': ('name', 'nombre'),
    'code': ('internal reference', 'referencia interna'),
    'location': ('location', 'ubicacion', 'ubicación'),
    'lot': ('lot/serial', 'lote/serie'),
    'qty': ('counted quantity', 'cantidad contada'),
}
REQUIRED_COLUMNS = {
    'name': 'Name',
    'code': 'Internal Reference',
    'location': 'Location',
    'qty': 'Counted Quantity',
}  # the Lot/Serial column is optional


class StockInventoryImportWizard(models.TransientModel):
    _name = 'ktm.stock.inventory.importer.wizard'
    _description = 'Physical Inventory File Validation'

    company_id = fields.Many2one(
        'res.company', required=True, readonly=True,
        default=lambda self: self.env.company,
    )
    file = fields.Binary(string='Excel File (.xlsx)', attachment=False)
    filename = fields.Char()
    state = fields.Selection(
        [('draft', 'Draft'), ('error', 'With Errors'), ('valid', 'Valid')],
        default='draft', readonly=True,
    )
    rows_checked = fields.Integer(readonly=True)
    result_message = fields.Char(readonly=True)
    error_ids = fields.One2many(
        'ktm.stock.inventory.importer.error', 'wizard_id',
        string='Errors', readonly=True,
    )

    # ------------------------------------------------------------------
    # Entry point
    # ------------------------------------------------------------------

    def action_validate(self):
        """Validate the uploaded file and reopen this wizard with the result.

        Nothing is written to stock: the check is read-only on purpose.
        """
        self.ensure_one()
        if not self.file:
            raise UserError(_("Please upload an Excel file first."))
        rows, errors = self._parse_xlsx(base64.b64decode(self.file))
        if not errors:
            errors = self._validate_rows(rows)
        errors.sort(key=lambda error: error['row_number'])
        self.write({
            'error_ids': [fields.Command.clear()] + [
                fields.Command.create(error) for error in errors
            ],
            'rows_checked': len(rows),
            'state': 'error' if errors else 'valid',
            'result_message': (
                _("%s error(s) found in the file.", len(errors)) if errors
                else _("The file is valid: %s row(s) checked, no errors.", len(rows))
            ),
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            # The widest size an action can request; the error list needs it.
            'context': {**self.env.context, 'dialog_size': 'extra-large'},
        }

    def action_open_grouped_errors(self):
        """Open the errors as a native list, grouped by message, in a new tab.

        A dialog would replace this wizard, so the list opens in a new
        browser tab. The wizard id travels in the URL path as ``active_id``;
        the action's domain reads it from the context.
        """
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': '/odoo/%s/action-%s' % (
                self.id,
                'ktm_stock_inventory_importer.'
                'ktm_stock_inventory_importer_error_action',
            ),
            'target': 'new',
        }

    # ------------------------------------------------------------------
    # Parsing
    # ------------------------------------------------------------------

    @api.model
    def _normalize_cell(self, value):
        """Return a cell as trimmed text; 1234.0 read from Excel becomes '1234'."""
        if value is None:
            return ''
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()

    @api.model
    def _parse_xlsx(self, content):
        """Read the first sheet of an xlsx file.

        :return: ``(rows, errors)``. ``rows`` is a list of dicts with the keys
            ``row_number``, ``name``, ``code``, ``location``, ``lot`` (all
            normalized text) and ``qty`` (the raw cell value). ``errors`` is
            non empty only when a required header is missing.
        :raises UserError: when the content is not a readable xlsx file.
        """
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                raw_rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        except Exception as exc:  # openpyxl raises many unrelated types
            raise UserError(_(
                "The file could not be read as an Excel (.xlsx) file: %s", exc,
            )) from exc
        if not raw_rows:
            raise UserError(_("The file is empty."))

        columns = self._map_header(raw_rows[0])
        missing = [
            label for key, label in REQUIRED_COLUMNS.items()
            if key not in columns
        ]
        if missing:
            return [], [{
                'row_number': 1,
                'field_name': _("Header"),
                'message': _(
                    "Missing required column(s): %s.", ", ".join(missing)),
            }]

        rows = []
        for index, raw in enumerate(raw_rows[1:], start=2):
            row = {
                key: raw[position] if position < len(raw) else None
                for key, position in columns.items()
            }
            if all(self._normalize_cell(value) == '' for value in row.values()):
                continue
            rows.append({
                'row_number': index,
                'name': self._normalize_cell(row.get('name')),
                'code': self._normalize_cell(row.get('code')),
                'location': self._normalize_cell(row.get('location')),
                'lot': self._normalize_cell(row.get('lot')),
                'qty': row.get('qty'),
            })
        return rows, []

    @api.model
    def _map_header(self, header_row):
        """Return ``{column key: position}`` for the recognized headers."""
        columns = {}
        for position, cell in enumerate(header_row):
            text = self._normalize_cell(cell).casefold()
            for key, aliases in HEADER_ALIASES.items():
                if text in aliases and key not in columns:
                    columns[key] = position
        return columns

    # ------------------------------------------------------------------
    # Validation (takes parsed rows, so a future load step can reuse it)
    # ------------------------------------------------------------------

    def _validate_rows(self, rows):
        """Validate parsed rows; return a list of error value dicts."""
        self.ensure_one()
        products = self._resolve_products(rows)
        locations = self._resolve_locations(rows)
        errors = []
        resolved = {}
        for row in rows:
            product, product_errors = products[row['row_number']]
            location, location_errors = locations[row['row_number']]
            resolved[row['row_number']] = (product, location)
            row_errors = product_errors + location_errors
            if product and not product_errors:
                row_errors += self._check_product_rules(product, row)
            row_errors += self._check_quantity(row)
            errors += [self._error_vals(row, *error) for error in row_errors]
        errors += self._check_duplicates(rows, resolved)
        return errors

    @api.model
    def _error_vals(self, row, field_name, message):
        return {
            'row_number': row['row_number'],
            'field_name': field_name,
            'message': message,
            'product_name': row['name'],
            'product_code': row['code'],
            'location_name': row['location'],
            'lot_name': row['lot'],
        }

    def _company_domain(self):
        return [('company_id', 'in', [False, self.company_id.id])]

    def _resolve_products(self, rows):
        """Match rows to products by internal reference AND name.

        Two products can share a code, hence both must match. One search for
        all the codes, then the name is compared in Python (exact and case
        insensitive, without ``=ilike`` wildcard surprises on ``%`` or ``_``).

        :return: ``{row_number: (product or None, [(field, message)])}``
        """
        codes = {row['code'] for row in rows if row['code']}
        by_code = defaultdict(list)
        if codes:
            candidates = self.env['product.product'].search(
                [('default_code', 'in', list(codes))] + self._company_domain())
            for product in candidates:
                by_code[product.default_code].append(product)
        result = {}
        for row in rows:
            product, errors = None, []
            if not row['name'] or not row['code']:
                errors.append((_("Product"), _(
                    "Product name and internal reference are both required.")))
            else:
                matches = [
                    product for product in by_code.get(row['code'], [])
                    if (product.name or '').strip().casefold()
                    == row['name'].casefold()
                ]
                if not matches:
                    errors.append((_("Product"), _(
                        "Product not found: no product matches both the "
                        "name and the internal reference.")))
                elif len(matches) > 1:
                    errors.append((_("Product"), _(
                        "Product is ambiguous: %s products match the name "
                        "and the internal reference.", len(matches))))
                else:
                    product = matches[0]
            result[row['row_number']] = (product, errors)
        return result

    def _resolve_locations(self, rows):
        """Match rows to internal locations (full name first, then name).

        :return: ``{row_number: (location or None, [(field, message)])}``
        """
        wanted = {row['location'].casefold() for row in rows if row['location']}
        by_complete, by_name = defaultdict(list), defaultdict(list)
        if wanted:
            for location in self.env['stock.location'].search(
                    self._company_domain()):
                by_complete[(location.complete_name or '').casefold()].append(location)
                by_name[(location.name or '').casefold()].append(location)
        result = {}
        for row in rows:
            location, errors = None, []
            key = row['location'].casefold()
            if not key:
                errors.append((_("Location"), _("Location is required.")))
            else:
                matches = by_complete.get(key) or by_name.get(key) or []
                internal = [loc for loc in matches if loc.usage == 'internal']
                if not matches:
                    errors.append((_("Location"), _("Location not found.")))
                elif not internal:
                    errors.append((_("Location"), _(
                        "Location is not an internal location.")))
                elif len(internal) > 1:
                    errors.append((_("Location"), _(
                        "Location is ambiguous: %s internal locations match.",
                        len(internal))))
                else:
                    location = internal[0]
            result[row['row_number']] = (location, errors)
        return result

    @api.model
    def _check_product_rules(self, product, row):
        errors = []
        if not product.is_storable:
            errors.append((_("Product"), _(
                "Product is not storable (Track Inventory is disabled).")))
        # A zero count means nothing is there, so there is no lot to name.
        elif (product.tracking in ('lot', 'serial') and not row['lot']
              and self._parse_quantity(row['qty'])[0] != 0):
            errors.append((_("Lot/Serial"), _(
                "A lot/serial number is required for a product tracked "
                "by lots or serial numbers.")))
        return errors

    @api.model
    def _parse_quantity(self, value):
        """Return ``(quantity, error message)``; quantity is None on error."""
        if value is None or (isinstance(value, str) and not value.strip()):
            return None, _("Counted quantity is required.")
        if isinstance(value, str) and ',' in value:
            return None, _(
                "Counted quantity contains a comma; use a numeric cell or "
                "a period as decimal separator.")
        try:
            quantity = float(value.strip()) \
                if isinstance(value, str) else float(value)
        except (TypeError, ValueError):
            return None, _("Counted quantity is not a number.")
        if quantity != quantity or quantity in (float('inf'), float('-inf')):
            return None, _("Counted quantity is not a number.")
        if quantity < 0:
            return None, _("Counted quantity cannot be negative.")
        return quantity, None

    @api.model
    def _check_quantity(self, row):
        _quantity, error = self._parse_quantity(row['qty'])
        return [(_("Counted Quantity"), error)] if error else []

    def _check_duplicates(self, rows, resolved):
        """Report each row repeating an earlier (product, location, lot)."""
        seen = defaultdict(list)
        errors = []
        for row in rows:
            product, location = resolved[row['row_number']]
            if not (row['name'] or row['code'] or row['location']):
                continue
            key = (
                product.id if product
                else ('raw', row['name'].casefold(), row['code'].casefold()),
                location.id if location else ('raw', row['location'].casefold()),
                row['lot'],
            )
            if seen[key]:
                errors.append(self._error_vals(row, _("Row"), _(
                    "Duplicate line: same product, location and lot as "
                    "row(s) %s.", ", ".join(map(str, seen[key])))))
            seen[key].append(row['row_number'])
        return errors


class StockInventoryImportError(models.TransientModel):
    _name = 'ktm.stock.inventory.importer.error'
    _description = 'Physical Inventory File Validation Error'
    _order = 'row_number, id'

    wizard_id = fields.Many2one(
        'ktm.stock.inventory.importer.wizard',
        required=True, ondelete='cascade',
    )
    # Not summed: a total of row numbers means nothing, and it would squeeze
    # the group name in the grouped list.
    row_number = fields.Integer(string='Row', readonly=True, aggregator=False)
    field_name = fields.Char(string='Column', readonly=True)
    product_name = fields.Char(string='Name', readonly=True)
    product_code = fields.Char(string='Internal Reference', readonly=True)
    location_name = fields.Char(string='Location', readonly=True)
    lot_name = fields.Char(string='Lot/Serial', readonly=True)
    message = fields.Char(readonly=True)
